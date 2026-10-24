"""Multi-format loader for the programme's shared legal corpus."""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from legal_rag.domain.entities import Document
from legal_rag.infrastructure.errors import DocumentSourceError

SUPPORTED_SUFFIXES = frozenset({".md", ".markdown", ".docx", ".pdf"})
_COUNTERPARTIES = {
    "northwind": "Northwind Logistics",
    "vertex": "Vertex Retail",
    "halcyon": "Halcyon Analytics",
}
_DOC_TYPES = (
    (r"amendment\s+(?:no\.?\s*)?\d", "amendment"),
    (r"schedules?\s+to", "schedule"),
    (r"non-disclosure", "nda"),
    (r"master\s+services\s+agreement", "master_agreement"),
)
_HEADER_LABELS = (
    r"Agreement No\.", r"Amendment No\.", r"Document No\.", r"Amends",
    r"Forms part of", r"Counterparty", r"Provider", r"Effective Date",
    r"Document Version", r"Status",
)
_STOP_LABELS = "|".join(_HEADER_LABELS)
_ID_PATTERN = re.compile(r"\b([A-Z]{2,4}-\d{4}-\d{3}(?:-\d{2})?)\b")


@dataclass(frozen=True)
class LoadedDocument:
    path: Path
    text: str
    metadata: dict[str, Any]
    table_count: int

    def to_entity(self) -> Document:
        document_type = (
            "amendment" if self.metadata["doc_type"] == "amendment" else "contract"
        )
        return Document(
            name=self.metadata["source_doc"],
            text=self.text,
            document_type=document_type,
            metadata=self.metadata,
        )


def _flatten_table(header: list[str], rows: list[list[str]]) -> str:
    lines = ["[TABLE] " + " | ".join(value for value in header if value)]
    for row in rows:
        cells = [
            f"{header[index]}: {cell}" if index < len(header) and header[index] else cell
            for index, cell in enumerate(row)
            if cell
        ]
        if cells:
            lines.append("  " + " | ".join(cells))
    return "\n".join(lines)


def _read_markdown(path: Path) -> tuple[str, int]:
    lines = path.read_text(encoding="utf-8").splitlines()
    output: list[str] = []
    table_count = 0
    index = 0

    def is_separator(line: str) -> bool:
        stripped = line.strip()
        return bool(stripped) and set(stripped) <= set("|-: ") and "-" in stripped and "|" in stripped

    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            header = [cell.strip() for cell in line.strip().strip("|").split("|")]
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            output.append(_flatten_table(header, rows))
            table_count += 1
            continue
        output.append(re.sub(r"[*_`]{1,2}", "", line))
        index += 1
    return "\n".join(output), table_count


def _read_docx(path: Path) -> tuple[str, int]:
    try:
        import docx
    except ImportError as exc:
        raise DocumentSourceError("python-docx is required to load .docx files") from exc

    source = docx.Document(str(path))
    body = source.element.body
    paragraphs = iter(source.paragraphs)
    tables = iter(source.tables)
    parts: list[str] = []
    table_count = 0
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = next(paragraphs, None)
            if paragraph is not None:
                parts.append(paragraph.text)
        elif tag == "tbl":
            table = next(tables, None)
            if table is not None and table.rows:
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                parts.append(_flatten_table(rows[0], rows[1:]))
                table_count += 1
    return "\n".join(parts), table_count


def _read_pdf(path: Path) -> tuple[str, int]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise DocumentSourceError("pdfplumber is required to load .pdf files") from exc

    parts: list[str] = []
    table_count = 0
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            body = page
            if tables:
                for table in page.find_tables():
                    body = body.outside_bbox(table.bbox)
            parts.append(body.extract_text() or "")
            for table in tables:
                rows = [
                    [cell.strip().replace("\n", " ") if cell else "" for cell in row]
                    for row in table if row
                ]
                if len(rows) >= 2:
                    parts.append(_flatten_table(rows[0], rows[1:]))
                    table_count += 1
    return "\n".join(parts), table_count


def _extract_metadata(text: str, path: Path) -> dict[str, Any]:
    head = text[:2000]
    title = ""
    for line in text.splitlines():
        stripped = line.strip()
        if stripped and not stripped.startswith("[TABLE]") and not line.startswith("  "):
            title = stripped.lstrip("# ").strip()
            break

    def grab(label: str) -> str | None:
        match = re.search(
            rf"(?:{label})\s*:?\s*\**\s*(.+?)(?=\s*(?:{_STOP_LABELS})\s*:|\n|$)",
            head,
            re.IGNORECASE | re.DOTALL,
        )
        return match.group(1).strip().strip("*").strip() if match else None

    source_doc = None
    id_labels = (
        (r"Amendment No\.", r"Document No\.", r"Agreement No\.")
        if re.search(r"amendment\s+(?:no\.?\s*)?\d", title, re.IGNORECASE)
        else (r"Document No\.", r"Agreement No\.", r"Amendment No\.")
    )
    for label in id_labels:
        match = re.search(
            rf"(?im)^\s*{label}\s*:\s*\**\s*{_ID_PATTERN.pattern}",
            head,
        )
        if match:
            source_doc = match.group(1)
            break
    if source_doc is None:
        match = _ID_PATTERN.search(head)
        source_doc = match.group(1) if match else path.stem

    counterparty_value = grab("Counterparty") or ""
    counterparty = next(
        (label for key, label in _COUNTERPARTIES.items() if key in counterparty_value.lower()),
        None,
    )
    effective_date = None
    raw_date = grab("Effective Date")
    date_match = re.match(r"(\d{1,2}\s+\w+\s+\d{4}|\d{4}-\d{2}-\d{2})", raw_date or "")
    if date_match:
        for date_format in ("%d %B %Y", "%d %b %Y", "%Y-%m-%d"):
            try:
                effective_date = datetime.strptime(date_match.group(1), date_format).strftime("%Y-%m-%d")
                break
            except ValueError:
                continue

    doc_type = next(
        (kind for pattern, kind in _DOC_TYPES if re.search(pattern, title, re.IGNORECASE)),
        "unknown",
    )
    status = grab("Status") or "unknown"
    amends_raw = grab("Amends") or grab("Forms part of")
    amends_match = _ID_PATTERN.search(amends_raw or "")
    amends_doc = amends_match.group(1) if amends_match else None

    metadata = {
        "source_doc": source_doc,
        "counterparty": counterparty,
        "effective_date": effective_date,
        "doc_type": doc_type,
        "status": status,
        "is_controlling": "CURRENT" in status.upper() or "CONTROLLING" in status.upper(),
        "amends": amends_doc,
    }
    metadata["source_doc"] = metadata["source_doc"] or "unknown"
    metadata["counterparty"] = metadata["counterparty"] or "Unknown"
    metadata["effective_date"] = metadata["effective_date"] or "Unknown"
    return metadata


def load_document(path: str | Path) -> LoadedDocument:
    source_path = Path(path)
    readers = {
        ".md": _read_markdown,
        ".markdown": _read_markdown,
        ".docx": _read_docx,
        ".pdf": _read_pdf,
    }
    suffix = source_path.suffix.lower()
    if suffix not in readers:
        raise DocumentSourceError(f"unsupported document format: {suffix}")
    text, table_count = readers[suffix](source_path)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    metadata = _extract_metadata(text, source_path)
    if metadata["doc_type"] == "schedule" and table_count == 0:
        raise DocumentSourceError(
            "schedule document produced no tables; refusing ingestion because schedule facts would be lost"
        )
    return LoadedDocument(source_path, text, metadata, table_count)


class SharedCorpusDocumentSource:
    """Filesystem source supporting one chosen corpus format per directory."""

    def __init__(self, docs_dir: Path):
        self._docs_dir = Path(docs_dir)

    def iter_documents(self, docs_dir: Path | None = None):
        source_dir = Path(docs_dir) if docs_dir else self._docs_dir
        if not source_dir.exists():
            raise DocumentSourceError(f"docs directory not found: {source_dir}")
        paths = sorted(
            path for path in source_dir.iterdir()
            if path.suffix.lower() in SUPPORTED_SUFFIXES
        )
        if not paths:
            raise DocumentSourceError(f"no supported documents found in {source_dir}")
        documents = [load_document(path).to_entity() for path in paths]
        source_ids = [document.metadata["source_doc"] for document in documents]
        if len(source_ids) != len(set(source_ids)):
            raise DocumentSourceError(
                "directory contains duplicate source_doc IDs; ingest one format at a time"
            )
        yield from documents

    def write_document(self, name: str, text: str) -> Path:
        self._docs_dir.mkdir(parents=True, exist_ok=True)
        target = self._docs_dir / name
        target.write_text(text, encoding="utf-8")
        return target


MarkdownDocumentSource = SharedCorpusDocumentSource
